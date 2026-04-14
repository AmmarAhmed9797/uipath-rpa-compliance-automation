"""
Compliance Report Generator — RPA Process
Predict360 GRC Platform — 360factors
Author: Muhammad Ammar Ahmed

This script automates compliance report generation using the Predict360 API,
replicating the UiPath RPA workflow in Python for portability and version control.
"""

import requests
import json
import pandas as pd
from datetime import datetime, timedelta
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
import smtplib
from email.mime.multipart import MIMEMultipart
from email.mime.base import MIMEBase
from email.mime.text import MIMEText
from email import encoders
import logging
import os
from config.config_reader import ConfigReader

# Configure logging
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(levelname)s - %(message)s',
    handlers=[
        logging.FileHandler('logs/compliance_report.log'),
        logging.StreamHandler()
    ]
)
logger = logging.getLogger(__name__)


class ComplianceReportGenerator:
    """Automates compliance report generation from Predict360 GRC platform."""

    def __init__(self):
        self.config = ConfigReader()
        self.base_url = self.config.get('api.base_url')
        self.auth_token = None
        self.report_data = {}
        self.wb = Workbook()

    def authenticate(self):
        """Authenticate with Predict360 API and get auth token."""
        logger.info("Authenticating with Predict360 API...")
        try:
            response = requests.post(
                f"{self.base_url}/api/auth/login",
                json={
                    "username": self.config.get('api.username'),
                    "password": self.config.get('api.password')
                },
                timeout=30
            )
            response.raise_for_status()
            self.auth_token = response.json()['token']
            logger.info("Authentication successful.")
        except requests.exceptions.RequestException as e:
            logger.error(f"Authentication failed: {e}")
            raise

    def fetch_compliance_alerts(self, days_back=30):
        """Fetch compliance alerts from the last N days."""
        logger.info(f"Fetching compliance alerts for the last {days_back} days...")
        start_date = (datetime.now() - timedelta(days=days_back)).strftime('%Y-%m-%d')
        end_date = datetime.now().strftime('%Y-%m-%d')

        response = requests.get(
            f"{self.base_url}/api/v1/compliance/alerts",
            headers={"Authorization": f"Bearer {self.auth_token}"},
            params={"startDate": start_date, "endDate": end_date, "pageSize": 1000},
            timeout=30
        )
        response.raise_for_status()
        alerts = response.json().get('data', [])
        logger.info(f"Fetched {len(alerts)} compliance alerts.")
        self.report_data['alerts'] = alerts
        return alerts

    def fetch_regulatory_changes(self):
        """Fetch recent regulatory changes."""
        logger.info("Fetching regulatory changes...")
        response = requests.get(
            f"{self.base_url}/api/v1/regulatory-changes",
            headers={"Authorization": f"Bearer {self.auth_token}"},
            params={"status": "Active", "pageSize": 500},
            timeout=30
        )
        response.raise_for_status()
        changes = response.json().get('data', [])
        logger.info(f"Fetched {len(changes)} regulatory changes.")
        self.report_data['regulatory_changes'] = changes
        return changes

    def generate_excel_report(self, output_path=None):
        """Generate formatted Excel compliance report."""
        if not output_path:
            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
            output_path = f"reports/compliance_report_{timestamp}.xlsx"

        os.makedirs(os.path.dirname(output_path), exist_ok=True)
        logger.info(f"Generating Excel report: {output_path}")

        # Alerts sheet
        ws_alerts = self.wb.active
        ws_alerts.title = "Compliance Alerts"
        headers = ['ID', 'Title', 'Severity', 'Status', 'Due Date', 'Assigned To', 'Regulation']
        self._write_header_row(ws_alerts, headers)

        for row_idx, alert in enumerate(self.report_data.get('alerts', []), start=2):
            ws_alerts.append([
                alert.get('id'), alert.get('title'),
                alert.get('severity'), alert.get('status'),
                alert.get('dueDate'), alert.get('assignedTo'),
                alert.get('regulation')
            ])

        # Summary sheet
        ws_summary = self.wb.create_sheet("Summary")
        self._add_summary_sheet(ws_summary)

        self.wb.save(output_path)
        logger.info(f"Report saved: {output_path}")
        return output_path

    def _write_header_row(self, ws, headers):
        """Write styled header row to worksheet."""
        header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)
        for col, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center')

    def _add_summary_sheet(self, ws):
        """Add executive summary sheet."""
        alerts = self.report_data.get('alerts', [])
        ws['A1'] = 'Compliance Report Summary'
        ws['A1'].font = Font(bold=True, size=14)
        ws['A3'] = 'Total Alerts:'
        ws['B3'] = len(alerts)
        severity_counts = {}
        for alert in alerts:
            sev = alert.get('severity', 'Unknown')
            severity_counts[sev] = severity_counts.get(sev, 0) + 1
        row = 4
        for severity, count in severity_counts.items():
            ws[f'A{row}'] = f'{severity} Severity:'
            ws[f'B{row}'] = count
            row += 1
        ws[f'A{row + 1}'] = f'Generated: {datetime.now().strftime("%Y-%m-%d %H:%M:%S")}'

    def send_email_report(self, report_path, recipients):
        """Send generated report via email."""
        logger.info(f"Sending report to: {recipients}")
        msg = MIMEMultipart()
        msg['From'] = self.config.get('email.sender')
        msg['To'] = ', '.join(recipients)
        msg['Subject'] = f"Compliance Report — {datetime.now().strftime('%B %Y')}"
        body = "Please find attached the monthly compliance report generated by the RPA bot."
        msg.attach(MIMEText(body, 'plain'))
        with open(report_path, 'rb') as f:
            part = MIMEBase('application', 'octet-stream')
            part.set_payload(f.read())
            encoders.encode_base64(part)
            part.add_header('Content-Disposition', f'attachment; filename={os.path.basename(report_path)}')
            msg.attach(part)
        with smtplib.SMTP(self.config.get('email.smtp_host'), 587) as server:
            server.starttls()
            server.login(self.config.get('email.username'), self.config.get('email.password'))
            server.sendmail(self.config.get('email.sender'), recipients, msg.as_string())
        logger.info("Report emailed successfully.")

    def run(self):
        """Main RPA workflow execution."""
        logger.info("=== Starting Compliance Report Generator RPA ===")
        try:
            self.authenticate()
            self.fetch_compliance_alerts(days_back=30)
            self.fetch_regulatory_changes()
            report_path = self.generate_excel_report()
            recipients = self.config.get('email.recipients').split(',')
            self.send_email_report(report_path, recipients)
            logger.info("=== Compliance Report Generator RPA Completed Successfully ===")
        except Exception as e:
            logger.error(f"RPA process failed: {e}")
            raise


if __name__ == "__main__":
    bot = ComplianceReportGenerator()
    bot.run()
